"""Import/Export routes — Excel XLSX."""

import io
from datetime import datetime

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook

from config import EXCEL_COLS
from auth import get_current_user
from helpers import parse_dob, audit_log
import database

router = APIRouter()


@router.get("/api/detainees/export/xlsx")
async def export_xlsx(user: dict = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Can pham"
    ws.append([h for _, h in EXCEL_COLS])
    async for d in database.db.detainees.find({}).sort("personal_id", 1):
        row = []
        for k, _ in EXCEL_COLS:
            v = d.get(k, "")
            row.append(v if v is not None else "")
        ws.append(row)
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"can_pham_{datetime.utcnow().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )


@router.get("/api/detainees/template/xlsx")
async def template_xlsx(user: dict = Depends(get_current_user)):
    wb = Workbook()
    ws = wb.active
    ws.title = "Mau nhap"
    ws.append([h for _, h in EXCEL_COLS])
    ws.append(["", "Nguyễn Văn Mẫu", "male", "15/03/1990", "001090123456", "Hà Nội", "Số 1, Hà Nội", "Kinh", "Không", "A01", "Trộm cắp tài sản", "01/01/2026", ""])
    for col in ws.columns:
        letter = col[0].column_letter
        ws.column_dimensions[letter].width = 18
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="mau_import.xlsx"'},
    )


@router.post("/api/detainees/import/xlsx")
async def import_xlsx(file: UploadFile = File(...), request: Request = None, user: dict = Depends(get_current_user)):
    if not (file.filename or "").lower().endswith(".xlsx"):
        raise HTTPException(400, "Chỉ nhận file .xlsx")
    data = await file.read()
    wb = load_workbook(io.BytesIO(data), read_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, "File rỗng")
    header = [str(x or "").strip() for x in rows[0]]
    header_map = {h: idx for idx, h in enumerate(header)}
    inserted, errors = 0, []
    now = datetime.utcnow()
    for i, r in enumerate(rows[1:], start=2):
        if not r or not any(r):
            continue
        try:
            def get(col_key):
                _, label = next((c for c in EXCEL_COLS if c[0] == col_key), (None, None))
                idx = header_map.get(label) if label else None
                if idx is None or idx >= len(r):
                    return None
                v = r[idx]
                return str(v).strip() if v is not None else None
            full_name = get("full_name")
            if not full_name:
                errors.append(f"Dòng {i}: thiếu Họ và tên")
                continue
            dob = parse_dob(get("dob"))
            date_in = parse_dob(get("date_in"))
            personal_id = (get("personal_id") or "").strip()
            if not personal_id:
                errors.append(f"Dòng {i}: thiếu mã can phạm (personal_id)")
                continue
            doc = {
                "personal_id": personal_id,
                "full_name": full_name,
                "gender": (get("gender") or "male").lower(),
                "dob": dob,
                "cccd_number": get("cccd_number") or "",
                "hometown": get("hometown"),
                "address": get("address"),
                "ethnicity": get("ethnicity"),
                "religion": get("religion"),
                "cell_code": get("cell_code"),
                "charge": get("charge"),
                "date_in": date_in,
                "note": get("note"),
                "created_at": now,
                "updated_at": now,
                "created_by": user["username"],
            }
            try:
                await database.db.detainees.insert_one(doc)
                inserted += 1
            except Exception as e:
                if "duplicate key" in str(e):
                    errors.append(f"Dòng {i}: số định danh {personal_id} đã tồn tại")
                else:
                    errors.append(f"Dòng {i}: {e}")
        except Exception as e:
            errors.append(f"Dòng {i}: {e}")
    await audit_log(request, user, "import", "detainee", "", {"inserted": inserted, "errors": len(errors)})
    return {"inserted": inserted, "errors": errors}
